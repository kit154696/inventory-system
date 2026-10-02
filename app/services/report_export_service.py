"""Excel export for the annual fiscal-year report, replicating server.js's
ExcelJS-based /api/report-summary/export layout via openpyxl.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from database.models import CategoryRepo, ReportRepo, SettingsRepo

THIN = Border(left=Side(style="thin"), right=Side(style="thin"),
              top=Side(style="thin"), bottom=Side(style="thin"))
HEADER_FILL = PatternFill("solid", fgColor="FFD9E1F2")
COLUMN_WIDTHS = [7, 25, 36, 12, 14, 12, 12, 15, 20]
HEADERS = ["ลำดับ", "ประเภท", "รายการ", "หน่วยนับ", "ยอดยกมา", "รับเข้า", "จ่ายออก", "คงเหลือ", "หมายเหตุ"]


def export_annual_report_excel(dest_path: Path, fiscal_year_be: int, cat_code: str | None = None) -> Path:
    dest_path = Path(dest_path)
    rows, totals = ReportRepo.annual_summary(fiscal_year_be, cat_code)

    org_name = SettingsRepo.get("orgName") or "ระบบทะเบียนคุมวัสดุ"
    category_label = "รวมทุกประเภท"
    if cat_code:
        cat = CategoryRepo.get(cat_code)
        if cat:
            category_label = cat["name"]

    wb = Workbook()
    ws = wb.active
    ws.title = "รายงาน"

    for i, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.merge_cells("A1:I1")
    ws["A1"] = "แบบรายงานรับจ่ายพัสดุประจำปีงบประมาณ"
    ws["A1"].font = Font(name="TH SarabunPSK", size=24, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 31

    ws.merge_cells("A2:I2")
    ws["A2"] = org_name
    ws["A2"].font = Font(name="TH SarabunPSK", size=20, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 26

    ws["A3"] = "รายงานการสำรวจ"
    ws["A3"].font = Font(name="TH SarabunPSK", size=20, bold=True)
    ws["C3"] = category_label
    ws["C3"].font = Font(name="TH SarabunPSK", size=20, bold=True)
    ws.row_dimensions[3].height = 26

    header_row = 4
    ws.row_dimensions[header_row].height = 22
    for i, text in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=i, value=text)
        cell.font = Font(name="TH SarabunPSK", size=16, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
        cell.fill = HEADER_FILL

    start_row = header_row + 1
    for idx, item in enumerate(rows):
        r = start_row + idx
        ws.row_dimensions[r].height = 22
        values = [
            item["no"], item["cat_name"], item["name"], item["unit"],
            item["brought_forward"], item["received"], item["issued"],
            item["balance"], "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = Font(name="TH SarabunPSK", size=16)
            cell.border = THIN
            cell.alignment = Alignment(
                horizontal="center" if (col == 1 or col >= 5) else "left",
                vertical="center",
            )

    total_row = start_row + len(rows)
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    ws.row_dimensions[total_row].height = 22
    total_label_cell = ws.cell(row=total_row, column=1, value="รวมทั้งสิ้น")
    total_label_cell.font = Font(name="TH SarabunPSK", size=16, bold=True)
    total_label_cell.alignment = Alignment(horizontal="center", vertical="center")
    total_label_cell.border = THIN

    total_values = [
        (5, totals["brought_forward"]), (6, totals["received"]),
        (7, totals["issued"]), (8, totals["balance"]),
    ]
    for col, value in total_values:
        cell = ws.cell(row=total_row, column=col, value=value)
        cell.font = Font(name="TH SarabunPSK", size=16, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN
    ws.cell(row=total_row, column=9).border = THIN

    wb.save(dest_path)
    return dest_path
